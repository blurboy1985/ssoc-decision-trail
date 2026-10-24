"""Convert SingStat's SSOC 2024 spreadsheets into data/ssoc2024.json.

Usage: python tools/build_ssoc.py raw/ssoc2024_definitions.xlsx raw/ssoc_correspondence.xlsx data/ssoc2024.json
"""

import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

# SingStat bullets come through as mojibake / private-use glyphs; strip them.
JUNK = re.compile(r"[�\u0080-\u009f -]+")


def clean(value) -> str:
    if value is None:
        return ""
    return JUNK.sub(" ", str(value)).replace("\r", "").strip()


def lines(value) -> list[str]:
    out = []
    for line in clean(value).split("\n"):
        line = line.strip().lstrip("-•").strip()
        if line:
            out.append(line)
    return out


def parent_of(code: str, nodes: dict) -> str | None:
    # 1 -> 11 -> 111 -> 1111 -> 11111 (major, sub-major, minor, unit, occupation).
    # Group X skips levels (X -> X1000), so take the longest prefix that exists.
    for end in range(len(code) - 1, 0, -1):
        if code[:end] in nodes:
            return code[:end]
    return None


def load_nodes(path: str) -> dict:
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        if row and row[0] == "SSOC 2024":
            break
    nodes = {}
    for row in rows:
        code = clean(row[0])
        if not code or not re.fullmatch(r"[0-9X]{1,5}", code):
            continue
        nodes[code] = {
            "code": code,
            # "ENGINEERING PROFESSIONALS II (Cont'd)" -> continuation pages of the printed report
            "title": re.sub(r"\s*\(Cont\W*d\)", "", clean(row[1])),
            "level": len(code),
            "parent": None,
            "children": [],
            "def": clean(row[3]),
            "tasks": lines(row[4]),
            "notes": clean(row[5]),
            "examples": lines(row[6]),
            "elsewhere": lines(row[7]),
        }
    for code, node in nodes.items():
        node["parent"] = parent_of(code, nodes)
        if node["parent"]:
            nodes[node["parent"]]["children"].append(code)
    return nodes


def load_correspondence(path: str) -> dict[str, list[str]]:
    ws = openpyxl.load_workbook(path, read_only=True)["SSOC2020-2024"]
    mapping: dict[str, list[str]] = {}
    for row in ws.iter_rows(values_only=True):
        old, new = clean(row[1]), clean(row[2])
        if re.fullmatch(r"[0-9X]{5}", old) and re.fullmatch(r"[0-9X]{5}", new):
            mapping.setdefault(old, [])
            if new not in mapping[old]:
                mapping[old].append(new)
    return mapping


def main() -> None:
    definitions, correspondence, out = sys.argv[1:4]
    nodes = load_nodes(definitions)
    roots = [code for code, node in nodes.items() if node["level"] == 1]
    data = {
        "version": "SSOC 2024",
        "roots": roots,
        "nodes": nodes,
        "ssoc2020to2024": load_correspondence(correspondence),
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    by_level = {}
    for node in nodes.values():
        by_level[node["level"]] = by_level.get(node["level"], 0) + 1
    widest = max(nodes.values(), key=lambda n: len(n["children"]))
    print("roots:", roots)
    print("nodes per level:", dict(sorted(by_level.items())))
    print("widest node:", widest["code"], widest["title"], len(widest["children"]))
    print("non-leaf nodes without children:", [c for c, n in nodes.items() if n["level"] < 5 and not n["children"]])
    print("2020->2024 mappings:", len(data["ssoc2020to2024"]))


if __name__ == "__main__":
    main()
